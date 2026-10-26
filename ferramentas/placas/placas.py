"""Placas retangulares de Kirchhoff–Love: solução de Ritz e tabelas para publicação.

Carga uniforme, bordas apoiadas (A) ou engastadas (E), cantos impedidos de
levantar. Convenções, dedução e testes independentes: README.md e tests/.
Somente este arquivo precisa acompanhar o notebook no Google Colab.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal
from functools import lru_cache
from pathlib import Path
import hashlib
import importlib.metadata
import json
import platform
import zipfile

import numpy as np
import pandas as pd
from numpy.polynomial import Legendre
from scipy.optimize import minimize, minimize_scalar
from scipy.sparse import csc_matrix, diags, eye, kron
from scipy.sparse.linalg import spsolve

VERSAO = "1.3.0"
POISSON = 0.20
# Ordem das bordas: x=0, x=lx, y=0, y=ly. x é o menor vão.
TIPOS = {
    "1":  ("A", "A", "A", "A"),
    "2A": ("A", "A", "E", "A"),
    "2B": ("E", "A", "A", "A"),
    "3":  ("E", "A", "E", "A"),
    "4A": ("A", "A", "E", "E"),
    "4B": ("E", "E", "A", "A"),
    "5A": ("E", "A", "E", "E"),
    "5B": ("E", "E", "E", "A"),
    "6":  ("E", "E", "E", "E"),
}
COLUNAS = ("mu_x", "mu_y", "mu_x_neg", "mu_y_neg", "alpha_w")
BORDAS = ("x0", "x1", "y0", "y1")
CANTOS = ("00", "10", "11", "01")
COLUNAS_REACOES = tuple(f"k_{b}" for b in BORDAS)
COLUNAS_CANTOS = tuple(f"c_{c}" for c in CANTOS)
COLUNAS_VOLVENTES = tuple(f"mu_xy_{c}" for c in CANTOS)
FONTES = {
    "equacao": "https://mbd.pages.rwth-aachen.de/courses/cmm/content/script/chapters/linearelastics.html",
    "tipologia": "https://wwwp.feb.unesp.br/pbastos/concreto1/Lajes.pdf",
    "reacoes": "https://cam.uni-wuppertal.de/fileadmin/bauing/baumechanik/FEM_1/Vorlesungen/L11_A_Kirchhoff.pdf",
}


def _numero(nome, valor, minimo, maximo):
    if isinstance(valor, (bool, np.bool_)):
        raise ValueError(f"{nome} deve ser um número, não um booleano.")
    try:
        valor = float(valor)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{nome} deve ser numérico.") from exc
    if not np.isfinite(valor) or not minimo <= valor <= maximo:
        raise ValueError(f"{nome} deve estar entre {minimo} e {maximo}.")
    return valor


def _tipo(tipo):
    tipo = str(tipo).strip().upper()
    if tipo not in TIPOS:
        raise ValueError(f"Tipo {tipo!r} inválido. Escolha entre {', '.join(TIPOS)}.")
    return tipo


@lru_cache(maxsize=256)
def _base(ordem, esquerda, direita):
    """Legendre integrado duas vezes, ortonormal em segunda derivada em [0,1].

    As funções sempre anulam w nas extremidades. Os modos k>=2 também
    anulam w'. Nos dois primeiros modos, uma combinação exata impõe um
    engaste; dois engastes eliminam ambos. Não há penalização de vínculos.
    """
    funcoes = []
    for k in range(ordem + 2):
        p = (Legendre.basis(k) * np.sqrt(2 * k + 1)).integ(2) / 4
        p = p - Legendre([(p(1) + p(-1)) / 2, (p(1) - p(-1)) / 2])
        funcoes.append(p)
    if esquerda == direita == "E":
        return tuple(funcoes[2:ordem + 2])
    if esquerda == direita == "A":
        return tuple(funcoes[:ordem])
    sinal = 1 if esquerda == "E" else -1
    primeira = (funcoes[0] + sinal * np.sqrt(3) * funcoes[1]) / 2
    return (primeira, *funcoes[2:ordem + 1])


@lru_cache(maxsize=1024)
def _coeficientes_base(ordem, esquerda, direita, derivada):
    polinomios = [p.deriv(derivada) * 2**derivada for p in _base(ordem, esquerda, direita)]
    coef = np.zeros((max(len(p.coef) for p in polinomios), ordem))
    for j, p in enumerate(polinomios):
        coef[:len(p.coef), j] = p.coef
    return coef


def _avaliar_base(ordem, esquerda, direita, t, derivada=0):
    coef = _coeficientes_base(ordem, esquerda, direita, derivada)
    valores = np.polynomial.legendre.legval(2 * np.asarray(t) - 1, coef)
    return np.moveaxis(valores, 0, -1)


@lru_cache(maxsize=256)
def _matrizes(ordem, esquerda, direita):
    base = _base(ordem, esquerda, direita)
    grau = max(len(p.coef) for p in base)
    pesos = 1 / (2 * np.arange(grau) + 1)
    matrizes = []
    for derivada in (0, 1):
        c = np.zeros((grau, ordem))
        for j, p in enumerate(base):
            coef = p.deriv(derivada).coef * 2**derivada
            c[:len(coef), j] = coef
        g = c.T @ (pesos[:, None] * c)
        # Remove somente arredondamento em produtos que deveriam ser nulos.
        escala = np.sqrt(np.outer(np.diag(g), np.diag(g)))
        g[np.abs(g) < 1e-13 * escala] = 0
        matrizes.append(csc_matrix(g))
    carga = np.array([p.coef[0] for p in base])  # integral em [0,1]
    return base, *matrizes, carga


@dataclass
class SolucaoPlaca:
    tipo: str
    razao: float
    poisson: float
    ordem: int
    amplitudes: np.ndarray
    residuo: float

    def derivada(self, u, v, dx=0, dy=0, grade=False):
        """Derivadas de w_hat = D*w/(p*lx**4) em u=x/lx, v=y/ly."""
        u, v = np.asarray(u, dtype=float), np.asarray(v, dtype=float)
        if not (np.all(np.isfinite(u)) and np.all(np.isfinite(v))
                and np.all((0 <= u) & (u <= 1)) and np.all((0 <= v) & (v <= 1))):
            raise ValueError("Coordenadas normalizadas u e v devem estar em [0,1].")
        if dx not in range(5) or dy not in range(5):
            raise ValueError("Ordem de derivada deve estar entre 0 e 4.")
        a, b, c, d = TIPOS[self.tipo]
        bx = _avaliar_base(self.ordem, a, b, u, dx)
        by = _avaliar_base(self.ordem, c, d, v, dy)
        if grade:
            return bx @ self.amplitudes @ by.T
        return np.einsum("...i,ij,...j->...", bx, self.amplitudes, by)

    def campo(self, nome, u, v, grade=False):
        """w_hat ou momentos normalizados por p*lx² (sem o fator 100)."""
        if nome == "w":
            return self.derivada(u, v, grade=grade)
        if nome == "mx":
            return -(self.derivada(u, v, 2, 0, grade) + self.poisson / self.razao**2 * self.derivada(u, v, 0, 2, grade))
        if nome == "my":
            return -(self.derivada(u, v, 0, 2, grade) / self.razao**2 + self.poisson * self.derivada(u, v, 2, 0, grade))
        if nome == "mxy":
            return -(1 - self.poisson) / self.razao * self.derivada(u, v, 1, 1, grade)
        raise ValueError("Campo deve ser w, mx, my ou mxy.")

    def _gradiente(self, nome, u, v):
        d = self.derivada
        if nome == "w":
            return np.array([d(u, v, 1, 0), d(u, v, 0, 1)])
        if nome == "mx":
            return -np.array([d(u, v, 3, 0) + self.poisson / self.razao**2 * d(u, v, 1, 2),
                              d(u, v, 2, 1) + self.poisson / self.razao**2 * d(u, v, 0, 3)])
        if nome == "my":
            return -np.array([d(u, v, 1, 2) / self.razao**2 + self.poisson * d(u, v, 3, 0),
                              d(u, v, 0, 3) / self.razao**2 + self.poisson * d(u, v, 2, 1)])
        return -(1 - self.poisson) / self.razao * np.array([d(u, v, 2, 1), d(u, v, 1, 2)])

    def valores_fisicos(self, lx, espessura, E, carga, u=0.5, v=0.5):
        """SI consistente: lx e h [m], E [kN/m²], p [kN/m²]; w [m], m [kN m/m]."""
        lx = _numero("lx [m]", lx, 1e-6, 1e6)
        espessura = _numero("espessura [m]", espessura, 1e-9, 1e4)
        E = _numero("E [kN/m²]", E, 1e-9, 1e15)
        carga = _numero("carga [kN/m²]", carga, 0, 1e12)
        rigidez = E * espessura**3 / (12 * (1 - self.poisson**2))
        return {"w_m": self.campo("w", u, v) * carga * lx**4 / rigidez,
                **{f"{nome}_kNm_m": self.campo(nome, u, v) * carga * lx**2 for nome in ("mx", "my", "mxy")}}


def resolver_placa(tipo="1", razao=1.0, poisson=POISSON, ordem=16):
    """Resolve D*laplaciano²(w)=p por Ritz em base tensorial de ordem² termos.

    Esta chamada resolve uma ordem. Use gerar_tabelas para testar convergência.
    """
    tipo = _tipo(tipo)
    razao = _numero("ly/lx", razao, 1, 10)
    poisson = _numero("Poisson", poisson, -0.99, 0.499)
    if isinstance(ordem, bool) or not isinstance(ordem, (int, np.integer)) or not 4 <= ordem <= 48:
        raise ValueError("ordem deve ser um inteiro entre 4 e 48.")
    a, b, c, d = TIPOS[tipo]
    _, mx, gx, fx = _matrizes(ordem, a, b)
    _, my, gy, fy = _matrizes(ordem, c, d)
    identidade = eye(ordem, format="csc")
    # Para todas as bordas com w=0, integral(f''g)=-integral(f'g').
    # Os termos de Poisson cancelam na forma fraca normalizada por D.
    K = (razao * kron(identidade, my) + kron(mx, identidade) / razao**3
         + 2 / razao * kron(gx, gy)).tocsc()
    f = razao * np.outer(fx, fy).ravel()
    escala = diags(1 / np.sqrt(K.diagonal()), format="csc")
    coef = escala @ spsolve((escala @ K @ escala).tocsc(), escala @ f)
    if not np.all(np.isfinite(coef)):
        raise ArithmeticError("A solução linear não é finita; reveja a ordem e a geometria.")
    denominador = np.max(np.asarray(abs(K).sum(axis=1))) * np.max(abs(coef)) + np.max(abs(f))
    residuo = float(np.max(abs(K @ coef - f)) / denominador)
    if residuo > 1e-9:
        raise ArithmeticError(f"Resíduo do sistema linear excessivo: {residuo:.3g}.")
    return SolucaoPlaca(tipo, razao, poisson, int(ordem), coef.reshape(ordem, ordem), residuo)


def _maximo(solucao, nome, sinal=1):
    """Busca na superfície: grade, múltiplos máximos locais e refinamento contínuo."""
    from scipy.ndimage import maximum_filter
    t = np.linspace(0, 1, max(41, 2 * solucao.ordem + 1))
    valores = sinal * solucao.campo(nome, t, t, grade=True)
    candidatos = np.argwhere(valores >= maximum_filter(valores, size=3))
    candidatos = sorted(candidatos, key=lambda ij: valores[tuple(ij)], reverse=True)[:8]
    ij = np.unravel_index(np.argmax(valores), valores.shape)
    melhor = (float(valores[ij]), float(t[ij[0]]), float(t[ij[1]]))
    for i, j in candidatos:
        ponto = np.array([t[i], t[j]])
        # A escala 100 melhora o critério de parada, inclusive para flechas pequenas.
        opt = minimize(lambda z: -100 * sinal * float(solucao.campo(nome, *z)), ponto,
                       jac=lambda z: -100 * sinal * solucao._gradiente(nome, *z),
                       bounds=((0, 1), (0, 1)), method="L-BFGS-B",
                       options={"ftol": 1e-14, "gtol": 1e-9, "maxiter": 150})
        valor = sinal * float(solucao.campo(nome, *opt.x))
        if valor > melhor[0]:
            melhor = (valor, float(opt.x[0]), float(opt.x[1]))
    return melhor


def _negativo_borda(solucao, borda):
    normal_x = borda < 2
    nome = "mx" if normal_x else "my"
    fixo = float(borda % 2)
    def valor(t):
        return -float(solucao.campo(nome, fixo, t) if normal_x else solucao.campo(nome, t, fixo))
    t = np.linspace(0, 1, max(101, 4 * solucao.ordem + 1))
    v = np.array([valor(s) for s in t])
    i = int(np.argmax(v))
    melhor = (max(0.0, float(v[i])), float(t[i]))
    for j in range(1, len(t) - 1):
        if v[j] >= v[j - 1] and v[j] >= v[j + 1]:
            opt = minimize_scalar(lambda s: -valor(s), bounds=(t[j - 1], t[j + 1]),
                                  method="bounded", options={"xatol": 1e-11})
            if -opt.fun > melhor[0]:
                melhor = (-float(opt.fun), float(opt.x))
    return melhor


def extrair_coeficientes(solucao):
    """Momentos/flecha extremos e volventes assinados nos quatro vértices.

    A conversão em momentos de armadura fica para o cálculo manual da aula 05.
    """
    coef, pontos = {}, []
    for nome, coluna, fator in (("mx", "mu_x", 100), ("my", "mu_y", 100),
                                ("w", "alpha_w", 1200 * (1 - solucao.poisson**2))):
        valor, u, v = _maximo(solucao, nome)
        coef[coluna] = fator * max(0, valor)
        pontos.append({"grandeza": coluna, "valor": coef[coluna], "x_lx": u, "y_ly": v})
    negativos = [0.0] * 4
    for borda, apoio in enumerate(TIPOS[solucao.tipo]):
        if apoio == "E":
            valor, s = _negativo_borda(solucao, borda)
            negativos[borda] = 100 * valor
            pontos.append({"grandeza": ("mu_neg_x0", "mu_neg_x1", "mu_neg_y0", "mu_neg_y1")[borda],
                           "valor": 100 * valor, "x_lx": float(borda % 2) if borda < 2 else s,
                           "y_ly": s if borda < 2 else float(borda % 2)})
    coef["mu_x_neg"] = max(negativos[:2])
    coef["mu_y_neg"] = max(negativos[2:])
    coef.update(dict(zip(("mu_neg_x0", "mu_neg_x1", "mu_neg_y0", "mu_neg_y1"), negativos)))
    for canto, (u, v) in zip(CANTOS, ((0, 0), (1, 0), (1, 1), (0, 1))):
        engastado = TIPOS[solucao.tipo][u] == "E" or TIPOS[solucao.tipo][2 + v] == "E"
        mu_xy = 0.0 if engastado else 100 * float(solucao.campo("mxy", u, v))
        coef[f"mu_xy_{canto}"] = mu_xy
    return coef, pontos


def reacao_borda(solucao, borda, t):
    """Reação efetiva de Kirchhoff r/(p*lx), positiva para cima sobre a placa.

    t em [0,1] percorre a borda; forças concentradas nos cantos são separadas.
    Inclui a derivada tangencial do momento volvente (não apenas Q normal).
    """
    if borda not in BORDAS:
        raise ValueError(f"Borda deve ser uma de {BORDAS}.")
    i = BORDAS.index(borda)
    fixo, sinal = float(i % 2), (-1 if i % 2 == 0 else 1)
    d, lam, nu = solucao.derivada, solucao.razao, solucao.poisson
    if i < 2:
        return sinal * (d(fixo, t, 3, 0) + (2 - nu) / lam**2 * d(fixo, t, 1, 2))
    return sinal * (d(t, fixo, 0, 3) / lam**3 + (2 - nu) / lam * d(t, fixo, 2, 1))


def reacoes_elasticas(solucao):
    """Médias k=10*r_media/(p*lx), resultantes R/(p*lx²), cantos C/(p*lx²).

    Integração gaussiana exata para os polinômios da aproximação. Não se
    corrige artificialmente a soma para fechar o equilíbrio da solução.
    """
    z, pesos = np.polynomial.legendre.leggauss(solucao.ordem + 4)
    t, pesos = (z + 1) / 2, pesos / 2
    dados, resultantes, cantos = {}, [], []
    for i, borda in enumerate(BORDAS):
        media = float(pesos @ reacao_borda(solucao, borda, t))
        comprimento = solucao.razao if i < 2 else 1
        resultantes.append(media * comprimento)
        dados[f"k_{borda}"] = 10 * media
        dados[f"R_{borda}"] = media * comprimento
    for nome, (u, v) in zip(CANTOS, ((0, 0), (1, 0), (1, 1), (0, 1))):
        # Um engaste adjacente impõe w_nt=0 exatamente.
        engastado = TIPOS[solucao.tipo][u] == "E" or TIPOS[solucao.tipo][2 + v] == "E"
        c = 0.0 if engastado else 2 * (-1)**(u + v) * float(solucao.campo("mxy", u, v))
        cantos.append(c)
        dados[f"c_{nome}"] = c
    dados["soma_R"] = sum(resultantes)
    dados["soma_C"] = sum(cantos)
    dados["carga_total"] = solucao.razao
    dados["erro_equilibrio"] = abs((sum(resultantes) + sum(cantos)) / solucao.razao - 1)
    return dados


@dataclass
class TabelaPlaca:
    tipo: str
    poisson: float
    tabela: pd.DataFrame
    convergencia: pd.DataFrame
    extremos: pd.DataFrame
    solucoes: dict[float, SolucaoPlaca]
    tolerancia: float
    ordens: tuple[int, ...]


def razoes_intervalo(inicio=1.0, fim=2.0, passo=0.05):
    inicio = _numero("lambda inicial", inicio, 1, 10)
    fim = _numero("lambda final", fim, inicio, 10)
    passo = _numero("passo", passo, 0.001, 9)
    a, b, s = map(lambda x: Decimal(str(x)), (inicio, fim, passo))
    n = (b - a) / s
    if n != n.to_integral_value():
        raise ValueError("O intervalo (final - inicial) deve ser múltiplo do passo.")
    if n > 500:
        raise ValueError("Gere no máximo 501 relações de vãos por execução.")
    return [float(a + i * s) for i in range(int(n) + 1)]


def gerar_tabelas(tipos="3", inicio=1.0, fim=2.0, passo=0.05, poisson=POISSON,
                  ordens=(8, 12, 16, 20, 24, 28, 32, 36, 40), tolerancia=0.001, progresso=None):
    """Gera tabelas; exige duas comparações sucessivas dentro da tolerância.

    O erro informado é a mudança relativa entre ordens, não um limite rigoroso
    de erro. Cada coeficiente e cada engaste é verificado, além do resíduo linear.
    """
    if isinstance(tipos, str):
        tipos = list(TIPOS) if tipos.strip().upper() == "TODOS" else [tipos]
    tipos = list(dict.fromkeys(_tipo(t) for t in tipos))
    if not tipos:
        raise ValueError("Selecione ao menos um tipo de laje.")
    poisson = _numero("Poisson", poisson, -0.99, 0.499)
    tolerancia = _numero("tolerância relativa", tolerancia, 1e-7, 0.05)
    ordens = tuple(ordens)
    if (len(ordens) < 3 or any(isinstance(n, bool) or not isinstance(n, (int, np.integer)) or not 4 <= n <= 48 for n in ordens)
            or any(a >= b for a, b in zip(ordens, ordens[1:]))):
        raise ValueError("Informe ao menos três ordens inteiras crescentes entre 4 e 48.")
    razoes = razoes_intervalo(inicio, fim, passo)
    resultado = {}
    total, atual = len(tipos) * len(razoes), 0
    for tipo in tipos:
        linhas, historico, extremos, solucoes = [], [], [], {}
        for razao in razoes:
            anterior, sucessos = None, 0
            for ordem in ordens:
                solucao = resolver_placa(tipo, razao, poisson, ordem)
                coef, pontos = extrair_coeficientes(solucao)
                reacoes = reacoes_elasticas(solucao)
                vetor = np.array([*coef.values(), *(reacoes[k] for k in (*COLUNAS_REACOES, *COLUNAS_CANTOS))])
                erro = np.nan if anterior is None else float(np.max(abs(vetor - anterior) / np.maximum(np.maximum(abs(vetor), abs(anterior)), 1e-3)))
                sucessos = sucessos + 1 if np.isfinite(erro) and erro <= tolerancia and reacoes["erro_equilibrio"] <= tolerancia else 0
                historico.append({"lambda": razao, "ordem": ordem, **coef, **reacoes,
                                  "mudanca_relativa": erro, "residuo_linear": solucao.residuo})
                anterior = vetor
                if sucessos >= 2:
                    break
            convergiu = sucessos >= 2
            linhas.append({"lambda": razao, **coef, **reacoes, "ordem": ordem,
                           "mudanca_relativa": erro, "convergiu": convergiu})
            extremos.extend({"lambda": razao, **ponto} for ponto in pontos)
            solucoes[razao] = solucao
            atual += 1
            if progresso is not None:
                progresso(atual, total, tipo, razao, convergiu)
        resultado[tipo] = TabelaPlaca(tipo, poisson, pd.DataFrame(linhas), pd.DataFrame(historico),
                                     pd.DataFrame(extremos), solucoes, tolerancia, ordens)
    return resultado


def desenhar_vinculos(tipo, ax=None, compacto=False):
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle
    tipo = _tipo(tipo)
    if ax is None:
        _, ax = plt.subplots(figsize=(4, 4))
    ax.add_patch(Rectangle((0, 0), 1, 1.4, fill=False, linewidth=0.8, color="black"))
    for i, apoio in enumerate(TIPOS[tipo]):
        if i < 2:
            x = float(i)
            ax.plot([x, x], [0, 1.4], color="black", lw=2.5 if apoio == "E" else 1)
            for y in np.linspace(0.1, 1.3, 9):
                if apoio == "E":
                    dx = -0.075 if i == 0 else 0.075
                    ax.plot([x, x + dx], [y, y - 0.06], color="black", lw=0.7)
                else:
                    ax.plot(x + (-0.04 if i == 0 else 0.04), y, marker=">" if i == 0 else "<",
                            ms=5, markerfacecolor="white", markeredgecolor="black")
        else:
            y = 0 if i == 2 else 1.4
            ax.plot([0, 1], [y, y], color="black", lw=2.5 if apoio == "E" else 1)
            for x in np.linspace(0.1, 0.9, 7):
                if apoio == "E":
                    dy = -0.075 if i == 2 else 0.075
                    ax.plot([x, x + 0.06], [y, y + dy], color="black", lw=0.7)
                else:
                    ax.plot(x, y + (-0.04 if i == 2 else 0.04), marker="^" if i == 2 else "v",
                            ms=5, markerfacecolor="white", markeredgecolor="black")
    ax.text(0.5, 0.7, f"Tipo {tipo}\np uniforme", ha="center", va="center", fontsize=7 if compacto else 11)
    ax.annotate("", (1, -0.19), (0, -0.19), arrowprops={"arrowstyle": "<->"})
    ax.text(0.5, -0.35, "$l_x$" if compacto else "$l_x$ (direção x)", ha="center", fontsize=9 if compacto else 10)
    ax.annotate("", (1.2, 1.4), (1.2, 0), arrowprops={"arrowstyle": "<->"})
    ax.text(1.31, 0.7, "$l_y$" if compacto else "$l_y$ (direção y)", rotation=90, va="center", fontsize=9 if compacto else 10)
    ax.set(xlim=(-0.2, 1.5), ylim=(-0.5, 1.65), aspect="equal")
    ax.axis("off")
    return ax.figure


def plotar_campos(solucao, pontos=101):
    import matplotlib.pyplot as plt
    pontos = int(_numero("pontos", pontos, 21, 501))
    t = np.linspace(0, 1, pontos)
    fig, axs = plt.subplots(2, 2, figsize=(9, 8), constrained_layout=True)
    for ax, nome, titulo, fator in zip(axs.flat, ("w", "mx", "my", "mxy"),
            (r"$Dw/(p l_x^4)$", r"$100 m_x/(p l_x^2)$", r"$100 m_y/(p l_x^2)$", r"$100 m_{xy}/(p l_x^2)$"),
            (1, 100, 100, 100)):
        campo = fator * solucao.campo(nome, t, t, grade=True)
        im = ax.contourf(t, solucao.razao * t, campo.T, levels=24, cmap="viridis" if nome == "w" else "RdBu_r")
        ax.set(xlabel="$x/l_x$", ylabel="$y/l_x$", title=titulo, aspect="equal")
        fig.colorbar(im, ax=ax, shrink=0.85)
    fig.suptitle(f"Tipo {solucao.tipo} | ly/lx = {solucao.razao:g} | ν = {solucao.poisson:g}")
    return fig


def desenhar_tipos(tipos=None):
    """Prancha vetorial dos vínculos para capa do Excel e material de aula."""
    import matplotlib.pyplot as plt
    tipos = list(TIPOS) if tipos is None else [_tipo(t) for t in tipos]
    if not tipos:
        raise ValueError("Selecione ao menos um tipo.")
    linhas = (len(tipos) + 2) // 3
    fig, axs = plt.subplots(linhas, 3, figsize=(10, 3.6 * linhas + 1.2), squeeze=False)
    for ax, tipo in zip(axs.flat, tipos):
        desenhar_vinculos(tipo, ax)
        ax.set_title(" · ".join(f"{b}: {a}" for b, a in zip(BORDAS, TIPOS[tipo])), fontsize=9)
    for ax in list(axs.flat)[len(tipos):]:
        ax.axis("off")
    fig.suptitle("Tipos de lajes retangulares — condições de apoio", fontsize=16, weight="bold", y=.98)
    fig.text(.5, .94, "Triângulos: apoio simples (A)    |    Hachuras: engaste (E)", ha="center", fontsize=11)
    fig.text(.5, .025, "x: menor vão  •  λ = ly/lx ≥ 1  •  Carga uniforme  •  Cantos impedidos de levantar",
             ha="center", fontsize=10)
    fig.subplots_adjust(top=.89, bottom=.07, hspace=.24, wspace=.15)
    return fig


def _familias_tabelas(item):
    """Dados, rótulos e fórmulas compartilhados por notebook, Excel e figuras."""
    t = item.tabela
    return {
        "Momentos": (t[["lambda", *COLUNAS]], ["λ", "μx", "μy", "μ'x", "μ'y", "αw"],
            "m = μ p lx²/100; negativos em módulo. wmax = αw p lx⁴/(100 E h³)."),
        "Volventes": (t[["lambda", *COLUNAS_VOLVENTES]], ["λ", "μxy00", "μxy10", "μxy11", "μxy01"],
            "Momento volvente no vértice: mxy = μxy p lx²/100 [kN m/m se p em kN/m² e lx em m]."),
        "Reacoes": (t[["lambda", *COLUNAS_REACOES]], ["λ", "kx0", "kx1", "ky0", "ky1"],
            "Reação média por borda: r = k p lx/10. Valores elásticos de borda, sem as forças "
            "concentradas de canto; a soma das bordas supera a carga total, a favor das vigas."),
    }


MOMENTOS_BARES = {
    "mu_x": ("μx", "Momento positivo na direção x: mx = μx p lx²/100."),
    "mu_y": ("μy", "Momento positivo na direção y: my = μy p lx²/100."),
    "mu_x_neg": ("μ'x", "Momento negativo nos engastes das bordas x0/x1, em módulo: m'x = μ'x p lx²/100. "
                        "Traço: tipo sem engaste nessa direção."),
    "mu_y_neg": ("μ'y", "Momento negativo nos engastes das bordas y0/y1, em módulo: m'y = μ'y p lx²/100. "
                        "Traço: tipo sem engaste nessa direção."),
    "mu_xy": ("μxy", "Momento volvente nos cantos com duas bordas apoiadas, em módulo: mxy = μxy p lx²/100. "
                     "Traço: tipo sem esse canto. Sinal de cada canto na aba do tipo."),
    "alpha_w": ("αw", "Flecha máxima: wmax = αw p lx⁴/(100 E h³)."),
    "k_x": ("kx", "Reação média nas bordas apoiadas x0/x1 (comprimento ly): r = kx p lx/10. "
                  "Traço: tipo sem borda apoiada nessa direção."),
    "k_x_neg": ("k'x", "Reação média nas bordas engastadas x0/x1 (comprimento ly): r = k'x p lx/10. "
                       "Traço: tipo sem engaste nessa direção."),
    "k_y": ("ky", "Reação média nas bordas apoiadas y0/y1 (comprimento lx): r = ky p lx/10. "
                  "Traço: tipo sem borda apoiada nessa direção."),
    "k_y_neg": ("k'y", "Reação média nas bordas engastadas y0/y1 (comprimento lx): r = k'y p lx/10. "
                       "Traço: tipo sem engaste nessa direção."),
}
# Reações de Bares: bordas da direção e tipo de vínculo procurado.
REACOES_BARES = {"k_x": (("x0", "x1"), "A"), "k_x_neg": (("x0", "x1"), "E"),
                 "k_y": (("y0", "y1"), "A"), "k_y_neg": (("y0", "y1"), "E")}
_CANTOS_UV = ((0, 0), (1, 0), (1, 1), (0, 1))
ABAS_BARES = {"mu_x": "Bares mx", "mu_y": "Bares my", "mu_x_neg": "Bares mx neg",
              "mu_y_neg": "Bares my neg", "mu_xy": "Bares mxy", "alpha_w": "Bares aw",
              "k_x": "Bares kx", "k_x_neg": "Bares kx eng", "k_y": "Bares ky", "k_y_neg": "Bares ky eng"}


def tabelas_por_momento(tabelas):
    """Formato de Bares: uma tabela por coeficiente, λ nas linhas e tipos nas colunas.

    Inclui momentos, flecha e reações médias de borda. NaN quando o coeficiente
    não existe no tipo: μ' e k' sem engaste na direção, k sem borda apoiada e
    μxy sem canto com duas bordas apoiadas. Retorna {coluna: (dados, rótulos, símbolo, legenda)}.
    """
    resultado = {}
    for coluna, (simbolo, legenda) in MOMENTOS_BARES.items():
        series = {}
        for tipo, item in tabelas.items():
            t = item.tabela.set_index("lambda")
            apoios = TIPOS[tipo]
            if coluna in REACOES_BARES:
                bordas, vinculo = REACOES_BARES[coluna]
                selecao = [f"k_{b}" for b in bordas if apoios[BORDAS.index(b)] == vinculo]
                existe = bool(selecao)
                valores = t[selecao].max(axis=1) if existe else t[f"k_{bordas[0]}"]
            elif coluna == "mu_xy":
                valores = t[list(COLUNAS_VOLVENTES)].abs().max(axis=1)
                existe = any(apoios[u] == apoios[2 + v] == "A" for u, v in _CANTOS_UV)
            else:
                valores = t[coluna]
                existe = {"mu_x_neg": "E" in apoios[:2], "mu_y_neg": "E" in apoios[2:]}.get(coluna, True)
            series[tipo] = valores if existe else valores * np.nan
        dados = pd.DataFrame(series).rename_axis("lambda").reset_index()
        resultado[coluna] = (dados, ["λ", *series], simbolo, legenda)
    return resultado


def _virgula(x, casas=2):
    return f"{x:.{casas}f}".replace(".", ",")


def _texto_coef(x, decimal=","):
    return "—" if pd.isna(x) else f"{x:.4f}".replace(".", decimal)


def _figura_bares(dados, rotulos, simbolo, legenda, poisson):
    import matplotlib.pyplot as plt
    import textwrap
    fig = plt.figure(figsize=(8.3, max(6, 2.6 + 0.24 * len(dados))))
    fig.text(0.06, 0.95, f"Coeficientes {simbolo} — ν = {_virgula(poisson)}", fontsize=14, weight="bold")
    fig.text(0.06, 0.915, "Carga uniforme • λ = ly/lx ≥ 1 • x: menor vão • tipos na prancha Tipos_de_lajes",
             fontsize=9)
    ax = fig.add_axes((0.06, 0.12, 0.88, 0.77))
    ax.axis("off")
    corpo = [[f"{row[0]:.2f}".replace(".", ","), *(_texto_coef(x) for x in row[1:])] for row in dados.to_numpy()]
    tabela = ax.table(cellText=corpo, colLabels=rotulos, cellLoc="center", bbox=(0, 0, 1, 1))
    tabela.auto_set_font_size(False)
    tabela.set_fontsize(8.5)
    for (i, j), cell in tabela.get_celld().items():
        cell.set_linewidth(0.4)
        if i == 0:
            cell.set_facecolor("#e5ebf0")
            cell.set_text_props(weight="bold")
    fig.text(0.07, 0.06, textwrap.fill(legenda, 94), fontsize=9)
    return fig


def exibir_tabelas_bares(tabelas):
    """Uma tabela por coeficiente, com os tipos lado a lado."""
    from IPython.display import Markdown, display
    for dados, rotulos, simbolo, legenda in tabelas_por_momento(tabelas).values():
        display(Markdown(f"### Tabela de {simbolo}\n{legenda}"))
        display(dados.set_axis(rotulos, axis=1).round(4).style.format(_texto_coef, subset=rotulos[1:]).hide(axis="index"))


def exibir_resultados(tabelas, lambda_visualizacao=1.5):
    """Interface curta do notebook; toda a apresentação fica neste módulo."""
    from IPython.display import Markdown, display
    import matplotlib.pyplot as plt
    for tipo, item in tabelas.items():
        display(Markdown(f"### Tipo {tipo} — Poisson {_virgula(item.poisson)}"))
        fig = desenhar_vinculos(tipo)
        display(fig)
        plt.close(fig)
        display(item.tabela[["lambda", *COLUNAS, "ordem", "mudanca_relativa", "convergiu"]].round(5))
        for nome, (dados, rotulos, legenda) in _familias_tabelas(item).items():
            if nome != "Momentos":
                display(Markdown(f"**{nome}** — {legenda}"))
                display(dados.set_axis(rotulos, axis=1).round(5))
        display(Markdown(f"Maior desequilíbrio relativo: **{100*item.tabela['erro_equilibrio'].max():.4f}%**."))
        nao = item.tabela.loc[~item.tabela["convergiu"], "lambda"].tolist()
        if nao:
            display(Markdown(f"**Convergência pendente para λ = {nao}. Aumente as ordens antes de exportar.**"))
        razao = min(item.solucoes, key=lambda r: abs(r - lambda_visualizacao))
        fig = plotar_campos(item.solucoes[razao])
        display(fig)
        plt.close(fig)
    display(Markdown("## Tabelas de Bares — um coeficiente por tabela"))
    exibir_tabelas_bares(tabelas)


def gerar_no_notebook(*args, **kwargs):
    """Barra de progresso para o Colab, mantendo callbacks fora do notebook."""
    from IPython.display import HTML, display
    painel = display(HTML("Preparando os cálculos…"), display_id=True)

    def atualizar(atual, total, tipo, razao, convergiu):
        estado = "convergência atendida" if convergiu else "convergência pendente"
        painel.update(HTML(f'<progress value="{atual}" max="{total}" style="width:70%"></progress>'
                          f'<p>{atual}/{total} — tipo {tipo}, λ = {razao:g}: {estado}.</p>'))

    return gerar_tabelas(*args, **kwargs, progresso=atualizar)


def baixar_resultados(arquivo):
    """Download no Colab; link para o ZIP quando executado no Jupyter local."""
    try:
        from google.colab import files
    except ImportError:
        import os
        from IPython.display import FileLink, display
        display(FileLink(os.path.relpath(arquivo), result_html_prefix="Baixar resultados: "))
    else:
        files.download(str(arquivo))


def _metadados(item):
    return {
        "programa": "Coeficientes de placas de Kirchhoff–Love", "versao": VERSAO,
        "tipo": item.tipo, "poisson": item.poisson,
        "bordas_x0_x1_y0_y1": list(TIPOS[item.tipo]),
        "carga": "uniforme, aplicada em toda a placa", "cantos": "levantamento impedido",
        "modelo": "placa delgada isotrópica, elástica linear, espessura constante",
        "metodo": "Rayleigh–Ritz com produtos de polinômios de Legendre integrados",
        "normalizacao_momentos": "m = mu * p * lx^2 / 100; mu_neg é módulo no engaste",
        "normalizacao_flecha": "w_max = alpha_w * p * lx^4 / (100 * E * h^3)",
        "normalizacao_volventes": "mxy_canto = mu_xy*p*lx^2/100; momento por unidade de largura",
        "uso_volventes": "momentos locais assinados; conversao Wood-Armer e uniformizacao realizadas manualmente na aula 05",
        "normalizacao_reacoes": "r_media = k*p*lx/10; R e c normalizados por p*lx^2",
        "sinal_reacoes": "positivo para cima sobre a placa; cantos negativos impedem levantamento",
        "equilibrio": "sum(R_bordas)+sum(C_cantos)=p*lx*ly; sem reescala artificial",
        "reacoes_exportadas": "apenas k das bordas; forcas de canto usadas so na verificacao do equilibrio",
        "rigidez": "D = E*h^3 / (12*(1-nu^2))",
        "extremos": "máximos procurados na superfície e nos engastes; posições na aba Extremos",
        "convergencia": "duas mudanças sucessivas <= tolerância; não é limite rigoroso de erro",
        "tolerancia_relativa": item.tolerancia, "ordens_disponiveis": list(item.ordens),
        "razoes": item.tabela["lambda"].tolist(),
        "nao_inclui": "bordas livres, fissuração, fluência, armaduras ou vigas deformáveis",
        "fontes": FONTES, "gerado_utc": datetime.now(timezone.utc).isoformat(),
        "python": platform.python_version(),
        "bibliotecas": {p: importlib.metadata.version(p) for p in ("numpy", "scipy", "pandas", "matplotlib", "openpyxl")},
        "sha256_modulo": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
    }


def _figura_tabela(item, familia="Momentos"):
    import matplotlib.pyplot as plt
    dados, rotulos, legenda = _familias_tabelas(item)[familia]
    fig = plt.figure(figsize=(8.3, max(6, 3.7 + 0.24 * len(dados))))
    ax = fig.add_axes((0.06, 0.81, 0.28, 0.17))
    desenhar_vinculos(item.tipo, ax, compacto=True)
    fig.text(0.39, 0.94, f"Placa tipo {item.tipo} — ν = {_virgula(item.poisson)}", fontsize=14, weight="bold")
    fig.text(0.39, 0.90, f"{familia} • Carga uniforme • λ = ly/lx ≥ 1", fontsize=10)
    fig.text(0.39, 0.86, "A: apoio simples; E: engaste\nCantos impedidos de levantar", fontsize=9)
    ax = fig.add_axes((0.06, 0.12, 0.88, 0.66))
    ax.axis("off")
    tabela = ax.table(cellText=[[f"{x:.4f}".replace(".", ",") for x in row] for row in dados.to_numpy()],
                      colLabels=rotulos, cellLoc="center", bbox=(0, 0, 1, 1))
    tabela.auto_set_font_size(False)
    tabela.set_fontsize(9)
    for (i, j), cell in tabela.get_celld().items():
        cell.set_linewidth(0.4)
        if i == 0:
            cell.set_facecolor("#e5ebf0")
            cell.set_text_props(weight="bold")
    import textwrap
    fig.text(0.07, 0.07, textwrap.fill(legenda, 94), fontsize=9)
    fig.text(0.07, 0.025, f"Tolerância entre ordens: {100*item.tolerancia:g}%. Metadados e histórico no Excel.", fontsize=8)
    return fig


def _exportar_excel(tabelas, destino, metadados):
    """Um livro: capa ilustrada, uma aba por tipo e auditoria consolidada."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.drawing.image import Image as ImagemExcel
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import Break

    with pd.ExcelWriter(destino / "Tabelas_lajes.xlsx", engine="openpyxl") as writer:
        capa = writer.book.create_sheet("Capa")
        capa["A1"] = "TABELAS DE LAJES RETANGULARES"
        capa["A1"].font = Font(size=20, bold=True, color="203C54")
        capa["A3"] = "Carga uniforme • Placa de Kirchhoff–Love • Cantos impedidos de levantar"
        capa["A4"] = "Triângulos: apoio simples (A). Hachuras: engaste (E). x acompanha o menor vão."
        capa["A5"] = ("Cada aba Tipo reúne momentos, flecha, volventes assinados, "
                      "e reações médias de borda (k).")
        capa["A6"] = "Reações: coeficientes k das bordas (r = k p lx/10); abas Bares kx, kx eng, ky e ky eng."
        capa["A8"] = "ABAS DE RESULTADOS (clique para abrir)"
        links = [(f"Tipo {tipo} — ν = {_virgula(item.poisson)} — {len(item.tabela)} relações de vãos", f"Tipo {tipo}")
                 for tipo, item in tabelas.items()]
        links += [(f"Tabela de Bares: {MOMENTOS_BARES[c][0]} para todos os tipos", aba) for c, aba in ABAS_BARES.items()]
        for linha, (texto, aba) in enumerate(links, 9):
            celula = capa.cell(linha, 1, texto)
            celula.hyperlink = f"#'{aba}'!A1"
            celula.style = "Hyperlink"
        desenho = ImagemExcel(str(destino / "Tipos_de_lajes.png"))
        desenho.width, desenho.height = 720, 720 * desenho.height / desenho.width
        capa.add_image(desenho, f"A{10 + len(links)}")
        capa.column_dimensions["A"].width = 110
        capa.sheet_view.showGridLines = False
        capa.print_area = f"A1:A{12 + len(links) + int(desenho.height / 20)}"
        capa.page_setup.orientation = "portrait"
        capa.page_setup.fitToWidth = 1
        capa.page_setup.fitToHeight = 1
        capa.sheet_properties.pageSetUpPr.fitToPage = True

        for tipo, item in tabelas.items():
            nome = f"Tipo {tipo}"
            folha = writer.book.create_sheet(nome)
            folha.sheet_view.showGridLines = False
            for linha, texto in enumerate((
                f"LAJE TIPO {tipo} | Poisson = {_virgula(item.poisson)} | carga uniforme",
                "x: menor vão; λ = ly/lx ≥ 1. A: apoio simples; E: engaste.",
                "Bordas x0, x1, y0, y1: " + ", ".join(TIPOS[tipo]),
                "Volventes: momentos por unidade de largura. Conversão para armaduras: cálculo manual.",
                "Consulte a aba Precisao para resultantes e verificação do equilíbrio.",
            ), 1):
                folha.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=6)
                folha.cell(linha, 1, texto)
            folha["A1"].font = Font(size=14, bold=True, color="203C54")
            prefixo = f"tipo_{tipo}_nu_{item.poisson:.2f}".replace(".", "p")
            desenho = ImagemExcel(str(destino / f"{prefixo}_apoios.png"))
            desenho.width, desenho.height = 220, 220 * desenho.height / desenho.width
            folha.add_image(desenho, "A7")
            linha = 7 + int(np.ceil(desenho.height / 20)) + 2
            primeira_linha = linha + 3
            for familia, (dados, rotulos, legenda) in _familias_tabelas(item).items():
                if familia != "Momentos":
                    folha.row_breaks.append(Break(id=linha - 1))
                folha.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=6)
                folha.cell(linha, 1, familia).font = Font(size=13, bold=True, color="203C54")
                folha.merge_cells(start_row=linha + 1, start_column=1, end_row=linha + 1, end_column=6)
                folha.cell(linha + 1, 1, legenda).alignment = Alignment(wrap_text=True, vertical="center")
                folha.row_dimensions[linha + 1].height = 32
                dados.set_axis(rotulos, axis=1).to_excel(writer, sheet_name=nome, startrow=linha + 2, index=False)
                for cell in folha[linha + 3][:len(rotulos)]:
                    cell.fill = PatternFill("solid", fgColor="203C54")
                    cell.font = Font(color="FFFFFF", bold=True)
                    cell.alignment = Alignment(horizontal="center")
                for row in folha.iter_rows(min_row=linha + 4, max_row=linha + 3 + len(dados), max_col=len(rotulos)):
                    for cell in row:
                        cell.number_format = "0.0000"
                        cell.alignment = Alignment(horizontal="center")
                linha += len(dados) + 7
            for col in range(1, 7):
                folha.column_dimensions[get_column_letter(col)].width = 20
            folha.freeze_panes = f"B{primeira_linha + 1}"
            folha.print_area = f"A1:F{linha - 3}"
            folha.print_title_rows = "1:3"
            folha.page_setup.orientation = "portrait"
            folha.page_setup.paperSize = folha.PAPERSIZE_A4
            folha.page_setup.fitToWidth = 1
            folha.page_setup.fitToHeight = 0
            folha.sheet_properties.pageSetUpPr.fitToPage = True

        poisson = next(iter(tabelas.values())).poisson
        for coluna, (dados, rotulos, simbolo, legenda) in tabelas_por_momento(tabelas).items():
            nome = ABAS_BARES[coluna]
            folha = writer.book.create_sheet(nome)
            folha.sheet_view.showGridLines = False
            folha.cell(1, 1, f"COEFICIENTES {simbolo} | Poisson = {_virgula(poisson)} | carga uniforme").font = \
                Font(size=14, bold=True, color="203C54")
            folha.cell(2, 1, "λ = ly/lx ≥ 1; x: menor vão. Colunas: tipos de laje (prancha na Capa).")
            folha.cell(3, 1, legenda)
            dados.set_axis(rotulos, axis=1).to_excel(writer, sheet_name=nome, startrow=4, index=False, na_rep="—")
            for cell in folha[5][:len(rotulos)]:
                cell.fill = PatternFill("solid", fgColor="203C54")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")
            for row in folha.iter_rows(min_row=6, max_row=5 + len(dados), max_col=len(rotulos)):
                for cell in row:
                    cell.number_format = "0.00" if cell.column == 1 else "0.0000"
                    cell.alignment = Alignment(horizontal="center")
            for col in range(1, len(rotulos) + 1):
                folha.column_dimensions[get_column_letter(col)].width = 11
            folha.freeze_panes = "B6"
            folha.print_title_rows = "1:5"
            folha.page_setup.orientation = "portrait"
            folha.page_setup.paperSize = folha.PAPERSIZE_A4
            folha.page_setup.fitToWidth = 1
            folha.page_setup.fitToHeight = 0
            folha.sheet_properties.pageSetUpPr.fitToPage = True

        for nome, atributo in (("Precisao", "tabela"), ("Convergencia", "convergencia"), ("Extremos", "extremos")):
            pd.concat([getattr(item, atributo).assign(tipo=tipo, poisson=item.poisson)
                       for tipo, item in tabelas.items()], ignore_index=True).to_excel(writer, sheet_name=nome, index=False)
        linhas = [(tipo, k, json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v)
                  for tipo, meta in metadados.items() for k, v in meta.items()]
        pd.DataFrame(linhas, columns=["Tipo", "Parametro", "Valor"]).to_excel(writer, sheet_name="Metadados", index=False)
        for nome in ("Precisao", "Convergencia", "Extremos", "Metadados"):
            folha = writer.sheets[nome]
            folha.freeze_panes = "B2"
            folha.auto_filter.ref = folha.dimensions
            for cell in folha[1]:
                cell.fill = PatternFill("solid", fgColor="203C54")
                cell.font = Font(color="FFFFFF", bold=True)
            for i in range(1, folha.max_column + 1):
                folha.column_dimensions[get_column_letter(i)].width = 24
            for row in folha.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = "0.00000000"
        writer.sheets["Metadados"].column_dimensions["C"].width = 100


def exportar_resultados(tabelas, pasta="resultados_placas"):
    """Excel único com capa e abas por tipo; desenhos, CSV, LaTeX, PDF/SVG e ZIP.

    Bloqueia linhas sem convergência e equilíbrio. Figuras com até 41 linhas
    por página. Preserva execuções anteriores em pastas com data e hora.
    """
    import matplotlib.pyplot as plt
    if not tabelas:
        raise ValueError("Não há tabelas para exportar.")
    for item in tabelas.values():
        if not item.tabela["convergiu"].all() or (item.tabela["erro_equilibrio"] > item.tolerancia).any():
            raise ValueError(f"Tipo {item.tipo}: há linhas sem convergência ou equilíbrio. Aumente ordens e gere novamente.")
    raiz = Path(pasta)
    raiz.mkdir(parents=True, exist_ok=True)
    destino = raiz / datetime.now(timezone.utc).strftime("placas_%Y%m%dT%H%M%S_%fZ")
    destino.mkdir()
    prancha = desenhar_tipos()
    for extensao in ("pdf", "svg", "png"):
        prancha.savefig(destino / f"Tipos_de_lajes.{extensao}", dpi=180, facecolor="white")
    plt.close(prancha)
    metadados = {}
    for tipo, item in tabelas.items():
        prefixo = f"tipo_{tipo}_nu_{item.poisson:.2f}".replace(".", "p")
        meta = _metadados(item)
        metadados[tipo] = meta
        (destino / f"{prefixo}_metadados.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
        fig = desenhar_vinculos(tipo)
        for extensao in ("png", "svg", "pdf"):
            fig.savefig(destino / f"{prefixo}_apoios.{extensao}", dpi=200, bbox_inches="tight", facecolor="white")
        plt.close(fig)
        for familia, (dados, rotulos, legenda) in _familias_tabelas(item).items():
            sufixo = "" if familia == "Momentos" else "_" + familia.lower()
            dados.to_csv(destino / f"{prefixo}{sufixo}.csv", sep=";", decimal=",", index=False, encoding="utf-8-sig")
            tex = ["% " + legenda, r"\begin{tabular}{" + "r" * len(rotulos) + "}", r"\hline",
                   " & ".join(rotulos) + r" \\", r"\hline"]
            tex += [" & ".join(f"{x:.4f}".replace(".", "{,}") for x in row) + r" \\" for row in dados.to_numpy()]
            tex += [r"\hline", r"\end{tabular}"]
            (destino / f"{prefixo}{sufixo}.tex").write_text("\n".join(tex) + "\n", encoding="utf-8")
            for pagina, inicio in enumerate(range(0, len(item.tabela), 41), 1):
                trecho = TabelaPlaca(item.tipo, item.poisson, item.tabela.iloc[inicio:inicio + 41], item.convergencia,
                                    item.extremos, item.solucoes, item.tolerancia, item.ordens)
                fig = _figura_tabela(trecho, familia)
                for extensao in ("pdf", "svg"):
                    fig.savefig(destino / f"{prefixo}{sufixo}_tabela_{pagina}.{extensao}", facecolor="white")
                plt.close(fig)
    poisson = next(iter(tabelas.values())).poisson
    for coluna, (dados, rotulos, simbolo, legenda) in tabelas_por_momento(tabelas).items():
        prefixo = f"bares_{coluna}_nu_{poisson:.2f}".replace(".", "p")
        dados.to_csv(destino / f"{prefixo}.csv", sep=";", decimal=",", index=False, encoding="utf-8-sig", na_rep="—")
        tex = ["% " + legenda, r"\begin{tabular}{" + "r" * len(rotulos) + "}", r"\hline",
               " & ".join(rotulos) + r" \\", r"\hline"]
        tex += [" & ".join([f"{row[0]:.2f}".replace(".", "{,}"), *(_texto_coef(x, "{,}") for x in row[1:])]) + r" \\"
                for row in dados.to_numpy()]
        tex += [r"\hline", r"\end{tabular}"]
        (destino / f"{prefixo}.tex").write_text("\n".join(tex) + "\n", encoding="utf-8")
        for pagina, inicio in enumerate(range(0, len(dados), 41), 1):
            fig = _figura_bares(dados.iloc[inicio:inicio + 41], rotulos, simbolo, legenda, poisson)
            for extensao in ("pdf", "svg"):
                fig.savefig(destino / f"{prefixo}_tabela_{pagina}.{extensao}", facecolor="white")
            plt.close(fig)
    _exportar_excel(tabelas, destino, metadados)
    (destino / "placas.py").write_bytes(Path(__file__).read_bytes())
    arquivo_zip = destino.with_suffix(".zip")
    with zipfile.ZipFile(arquivo_zip, "w", zipfile.ZIP_DEFLATED) as zipado:
        for arquivo in sorted(destino.iterdir()):
            zipado.write(arquivo, arquivo.name)
    return arquivo_zip
