"""Referências independentes (Navier e Lévy), vínculos e arquivos exportados.

Execute: python -m unittest discover -s ferramentas/placas/tests -v
"""
from pathlib import Path
import json
import sys
import tempfile
import unittest
import zipfile

import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from placas import (TIPOS, resolver_placa, extrair_coeficientes, gerar_tabelas,
                    exportar_resultados, razoes_intervalo, reacoes_elasticas,
                    tabelas_por_momento)


def navier(razao, nu, u, v, termos=201):
    """Série dupla de senos; independente das bases e matrizes de Ritz."""
    m = np.arange(1, 2 * termos, 2, dtype=float)[:, None]
    n = np.arange(1, 2 * termos, 2, dtype=float)[None, :]
    amplitude = 16 / (np.pi**6 * m * n * (m**2 + (n / razao)**2)**2)
    seno = np.sin(m * np.pi * u) * np.sin(n * np.pi * v)
    w = np.sum(amplitude * seno)
    mx = np.sum(np.pi**2 * (m**2 + nu * n**2 / razao**2) * amplitude * seno)
    my = np.sum(np.pi**2 * (n**2 / razao**2 + nu * m**2) * amplitude * seno)
    return np.array([w, mx, my])


def levy(apoios_x, razao, nu, u, v, termos=121):
    """Lévy: lados y apoiados; solução analítica da EDO de cada seno em y.

    f'''' - 2 k² f'' + k⁴ f = 4/(n*pi).
    Exponenciais escaladas evitam cancelamento de senh/cosh para n grande.
    """
    total = np.zeros(3)
    for n in range(1, 2 * termos, 2):
        k = n * np.pi / razao
        particular = 4 / (n * np.pi * k**4)

        def base(x, derivada):
            a, b = np.exp(-k * x), np.exp(-k * (1 - x))
            r = derivada
            return np.array([(-k)**r * a,
                (((-k)**r * x) + (r * (-k)**(r - 1) if r else 0)) * a,
                k**r * b, (k**r * (1 - x) - (r * k**(r - 1) if r else 0)) * b])

        matriz = np.stack([base(0, 0), base(0, 1 if apoios_x[0] == "E" else 2),
                            base(1, 0), base(1, 1 if apoios_x[1] == "E" else 2)])
        c = np.linalg.solve(matriz, [-particular, 0, -particular, 0])
        f = particular + base(u, 0) @ c
        f2 = base(u, 2) @ c
        total += np.sin(n * np.pi * v) * np.array([f, -f2 + nu * k*k*f, k*k*f - nu*f2])
    return total


class SolverTests(unittest.TestCase):
    def test_volventes_nos_cantos(self):
        for tipo in ("1", "3", "6"):
            s = resolver_placa(tipo, 1.5, ordem=28)
            coef, _ = extrair_coeficientes(s)
            r = reacoes_elasticas(s)
            for nome, (u, v) in zip(("00", "10", "11", "01"), ((0, 0), (1, 0), (1, 1), (0, 1))):
                self.assertAlmostEqual(coef[f"mu_xy_{nome}"], 100*float(s.campo("mxy", u, v)), places=10)
                self.assertAlmostEqual(r[f"c_{nome}"], 2*(-1)**(u+v)*coef[f"mu_xy_{nome}"]/100, places=12)
            if tipo == "1":
                self.assertLess(coef["mu_xy_00"], 0)
                self.assertGreater(coef["mu_xy_10"], 0)
            if tipo == "6":
                self.assertTrue(all(coef[f"mu_xy_{c}"] == 0 for c in ("00", "10", "11", "01")))

    def test_reacoes_navier_e_equilibrio(self):
        # Integração analítica da série dupla: independente da quadratura de Ritz.
        for lam in (1, 1.5, 2):
            for nu in (0, .2, .3):
                referencias = []
                for termos in (600, 1200):
                    m = np.arange(1, 2*termos, 2, dtype=float)[:, None]
                    n = np.arange(1, 2*termos, 2, dtype=float)[None, :]
                    amp = 16 / (np.pi**6 * m * n * (m*m + (n/lam)**2)**2)
                    rx = float(np.sum(amp * m*np.pi * ((m*np.pi)**2 + (2-nu)*(n*np.pi/lam)**2) * 2*lam/(n*np.pi)))
                    ry = float(np.sum(amp * n*np.pi/lam * ((n*np.pi/lam)**2 + (2-nu)*(m*np.pi)**2) * 2/(m*np.pi)))
                    canto = float(-2*(1-nu)/lam * np.sum(amp*m*n*np.pi**2))
                    referencias.append(np.array([rx, ry, canto]))
                # A integração do cortante tem cauda O(1/N); Richardson cancela esse termo.
                referencia = 2*referencias[1] - referencias[0]
                r = reacoes_elasticas(resolver_placa("1", lam, nu, ordem=36))
                np.testing.assert_allclose([r["R_x0"], r["R_y0"], r["c_00"]], referencia, rtol=2e-4)
                self.assertLess(r["c_00"], 0)
                self.assertGreater(r["soma_R"], lam)
                self.assertLess(r["erro_equilibrio"], .001)
        for tipo in TIPOS:
            r = reacoes_elasticas(resolver_placa(tipo, 1.5, ordem=36))
            self.assertLess(r["erro_equilibrio"], .001)
        r = reacoes_elasticas(resolver_placa("6", 1, ordem=36))
        self.assertEqual(r["soma_C"], 0)
        self.assertAlmostEqual(r["R_x0"], .25, delta=.0001)

    def test_volvente_navier(self):
        # Navier no vértice: mxy = (1-nu)*sum A*m*pi*n*pi/lambda (cossenos = 1).
        m = np.arange(1, 801, 2, dtype=float)[:, None]
        n = np.arange(1, 801, 2, dtype=float)[None, :]
        for lam in (1.0, 1.5, 2.0):
            amplitude = 16 / (np.pi**6 * m * n * (m**2 + (n / lam)**2)**2)
            esperado = 100 * 0.8 * np.sum(amplitude * m * np.pi * n * np.pi / lam)
            coef, _ = extrair_coeficientes(resolver_placa("1", lam, 0.2, 28))
            for c in ("00", "10", "11", "01"):
                self.assertAlmostEqual(abs(coef[f"mu_xy_{c}"]), esperado, delta=2e-3 * esperado)

    def test_tabelas_por_momento(self):
        tabelas = gerar_tabelas(["1", "2B", "6"], 1, 1.5, .5, ordens=(8, 12, 16))
        bares = tabelas_por_momento(tabelas)
        self.assertEqual(list(bares), ["mu_x", "mu_y", "mu_x_neg", "mu_y_neg", "mu_xy", "alpha_w",
                                       "k_x", "k_x_neg", "k_y", "k_y_neg"])
        for coluna, (dados, rotulos, _, _) in bares.items():
            self.assertEqual(rotulos, ["λ", "1", "2B", "6"])
            self.assertEqual(dados["lambda"].tolist(), [1.0, 1.5])
        for tipo in ("1", "2B", "6"):
            np.testing.assert_array_equal(bares["mu_x"][0][tipo], tabelas[tipo].tabela["mu_x"])
        # Traço quando o coeficiente não existe no tipo.
        self.assertTrue(bares["mu_x_neg"][0]["1"].isna().all())
        self.assertTrue(bares["mu_y_neg"][0]["2B"].isna().all())
        self.assertTrue(bares["mu_xy"][0]["6"].isna().all())
        np.testing.assert_array_equal(bares["mu_x_neg"][0]["2B"], tabelas["2B"].tabela["mu_x_neg"])
        t = tabelas["2B"].tabela
        np.testing.assert_array_equal(bares["mu_xy"][0]["2B"], t[["mu_xy_10", "mu_xy_11"]].abs().max(axis=1))
        # Reações: 2B tem x0 engastada e x1, y0, y1 apoiadas.
        np.testing.assert_array_equal(bares["k_x"][0]["2B"], t["k_x1"])
        np.testing.assert_array_equal(bares["k_x_neg"][0]["2B"], t["k_x0"])
        np.testing.assert_array_equal(bares["k_y"][0]["2B"], t[["k_y0", "k_y1"]].max(axis=1))
        self.assertTrue(bares["k_y_neg"][0]["2B"].isna().all())
        self.assertTrue(bares["k_x_neg"][0]["1"].isna().all())
        self.assertTrue(bares["k_x"][0]["6"].isna().all())
        np.testing.assert_allclose(bares["k_y_neg"][0]["6"], tabelas["6"].tabela["k_y0"])

    def test_navier_com_geometria_e_poisson_variados(self):
        for razao in (1, 1.35, 2):
            for nu in (0, 0.2, 0.3):
                with self.subTest(razao=razao, nu=nu):
                    s = resolver_placa("1", razao, nu, ordem=24)
                    esperado = navier(razao, nu, .37, .43)
                    obtido = [s.campo(k, .37, .43) for k in ("w", "mx", "my")]
                    np.testing.assert_allclose(obtido, esperado, rtol=2e-5, atol=1e-8)

    def test_levy_com_engastes(self):
        for tipo in ("2B", "4B"):
            for razao in (1, 1.5, 2):
                with self.subTest(tipo=tipo, razao=razao):
                    s = resolver_placa(tipo, razao, ordem=28)
                    for u, v in ((.31, .43), (.5, .5), (0, .5)):
                        esperado = levy(TIPOS[tipo][:2], razao, .2, u, v)
                        obtido = [s.campo(k, u, v) for k in ("w", "mx", "my")]
                        np.testing.assert_allclose(obtido, esperado, rtol=8e-5, atol=2e-7)

    def test_vinculos_de_todos_os_tipos(self):
        t = np.linspace(.1, .9, 13)
        for tipo, apoios in TIPOS.items():
            s = resolver_placa(tipo, 1.4, ordem=28)
            with self.subTest(tipo=tipo):
                for borda, apoio in enumerate(apoios):
                    u, v = (float(borda % 2), t) if borda < 2 else (t, float(borda % 2))
                    np.testing.assert_allclose(s.campo("w", u, v), 0, atol=2e-15)
                    if apoio == "E":
                        rotacao = s.derivada(u, v, 1 if borda < 2 else 0, 0 if borda < 2 else 1)
                        np.testing.assert_allclose(rotacao, 0, atol=2e-13)
                    else:
                        momento = s.campo("mx" if borda < 2 else "my", u, v)
                        np.testing.assert_allclose(momento, 0, atol=1e-4)

    def test_rotacao_quadrada_e_simetria(self):
        for a, b in (("2A", "2B"), ("4A", "4B"), ("5A", "5B")):
            sa, sb = resolver_placa(a, 1, ordem=20), resolver_placa(b, 1, ordem=20)
            for u, v in ((.2, .3), (.5, .7)):
                self.assertAlmostEqual(float(sa.campo("w", u, v)), float(sb.campo("w", v, u)), places=12)
                self.assertAlmostEqual(float(sa.campo("mx", u, v)), float(sb.campo("my", v, u)), places=11)
        s = resolver_placa("6", 1, ordem=24)
        # Referência clássica da placa quadrada engastada: w D/(p a^4) ≈ 0,00126.
        self.assertAlmostEqual(float(s.campo("w", .5, .5)), .0012653, delta=5e-7)
        self.assertAlmostEqual(float(s.campo("mx", .5, .5)), float(s.campo("my", .5, .5)), places=12)

    def test_extremos_e_unidades(self):
        s = resolver_placa("1", 1, ordem=24)
        coef, pontos = extrair_coeficientes(s)
        esperado = navier(1, .2, .5, .5)
        self.assertAlmostEqual(coef["mu_x"], 100 * esperado[1], delta=2e-5)
        self.assertAlmostEqual(coef["alpha_w"], 1200 * .96 * esperado[0], delta=1e-6)
        self.assertEqual(coef["mu_x_neg"], 0)
        self.assertTrue(all(abs(p["x_lx"] - .5) < 1e-5 for p in pontos))
        a = s.valores_fisicos(4, .12, 30e6, 5)
        b = s.valores_fisicos(8, .12, 30e6, 5)
        c = s.valores_fisicos(4, .24, 30e6, 5)
        dobro_carga = s.valores_fisicos(4, .12, 30e6, 10)
        dobro_E = s.valores_fisicos(4, .12, 60e6, 5)
        self.assertAlmostEqual(float(b["w_m"] / a["w_m"]), 16)
        self.assertAlmostEqual(float(b["mx_kNm_m"] / a["mx_kNm_m"]), 4)
        self.assertAlmostEqual(float(c["w_m"] / a["w_m"]), .125)
        self.assertAlmostEqual(float(dobro_carga["w_m"] / a["w_m"]), 2)
        self.assertAlmostEqual(float(dobro_carga["mx_kNm_m"] / a["mx_kNm_m"]), 2)
        self.assertAlmostEqual(float(dobro_E["w_m"] / a["w_m"]), .5)
        self.assertAlmostEqual(float(dobro_E["mx_kNm_m"] / a["mx_kNm_m"]), 1)

    def test_parametros_invalidos(self):
        for kwargs in ({"tipo": "7"}, {"razao": .9}, {"poisson": .5},
                       {"poisson": float("nan")}, {"ordem": True}, {"ordem": 3}):
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                resolver_placa(**kwargs)
        self.assertEqual(len(razoes_intervalo(1, 2, .05)), 21)
        self.assertEqual(razoes_intervalo(1, 1), [1])
        for args in ((1, 2, .3), (2, 1, .1), (1, 2, 0)):
            with self.assertRaises(ValueError):
                razoes_intervalo(*args)
        with self.assertRaises(ValueError):
            gerar_tabelas([], 1, 1)


class ExportacaoTests(unittest.TestCase):
    def test_convergencia_e_arquivos(self):
        tabelas = gerar_tabelas(["1", "3"], 1, 1.1, .1)
        item = tabelas["3"]
        self.assertTrue(item.tabela["convergiu"].all())
        for razao, grupo in item.convergencia.groupby("lambda"):
            self.assertGreaterEqual(len(grupo), 3)
            self.assertTrue((grupo["mudanca_relativa"].iloc[-2:] <= item.tolerancia).all())
        with tempfile.TemporaryDirectory() as temp:
            arquivo = exportar_resultados(tabelas, temp)
            with zipfile.ZipFile(arquivo) as z:
                nomes = z.namelist()
                self.assertEqual(sum(n.endswith(".xlsx") for n in nomes), 1)
                for extensao in (".csv", ".tex", ".pdf", ".svg", ".json"):
                    self.assertTrue(any(n.endswith(extensao) for n in nomes))
                meta = json.loads(z.read(next(n for n in nomes if n.startswith("tipo_3_") and n.endswith(".json"))))
                self.assertEqual(meta["poisson"], .2)
                self.assertEqual(meta["tipo"], "3")
                self.assertEqual(meta["bordas_x0_x1_y0_y1"], ["E", "A", "E", "A"])
            xlsx = next(arquivo.with_suffix("").glob("*.xlsx"))
            wb = load_workbook(xlsx, data_only=True)
            self.assertEqual(set(wb.sheetnames), {"Capa", "Tipo 1", "Tipo 3", "Precisao", "Convergencia", "Extremos", "Metadados",
                                                 "Bares mx", "Bares my", "Bares mx neg", "Bares my neg",
                                                 "Bares mxy", "Bares aw", "Bares kx", "Bares kx eng",
                                                 "Bares ky", "Bares ky eng"})
            bares = wb["Bares mx neg"]
            self.assertEqual([c.value for c in bares[5][:3]], ["λ", "1", "3"])
            self.assertEqual(bares.cell(6, 2).value, "—")
            self.assertAlmostEqual(bares.cell(6, 3).value, tabelas["3"].tabela.iloc[0]["mu_x_neg"])
            folha = wb["Tipo 3"]
            cabecalhos = [row[0].row for row in folha if row[0].value == "λ"]
            # Momentos, Volventes, Reacoes
            self.assertEqual(len(cabecalhos), 3)
            self.assertEqual(folha.cell(cabecalhos[0] + 1, 1).value, 1)
            self.assertAlmostEqual(folha.cell(cabecalhos[0] + 1, 2).value, item.tabela.iloc[0]["mu_x"])
            self.assertAlmostEqual(folha.cell(cabecalhos[1] + 1, 4).value, item.tabela.iloc[0]["mu_xy_11"])
            self.assertAlmostEqual(folha.cell(cabecalhos[2] + 1, 2).value, item.tabela.iloc[0]["k_x0"])
            self.assertEqual(len(folha._images), 1)
            self.assertEqual(len(wb["Capa"]._images), 1)
            wb.close()
            item.tabela.loc[0, "convergiu"] = False
            with self.assertRaisesRegex(ValueError, "sem convergência"):
                exportar_resultados(tabelas, temp)


if __name__ == "__main__":
    unittest.main()
